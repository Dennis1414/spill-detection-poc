"""Build the Word version of the demo report (report/SpillDetect-report.docx).

    python scripts/build_docx.py

Same data and the same selection of examples as the HTML report, without the
collapsible appendices. Boxes are drawn onto the images (amber = model's answer,
green = true spill) because Word cannot overlay them.
"""
from __future__ import annotations

import io
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))

import build_report as br  # noqa: E402
from spilldetect.images import load_resized  # noqa: E402

OUT = ROOT / "report" / "SpillDetect-report.docx"
AMBER = (224, 161, 0)
GREEN = (31, 138, 91)
INK = RGBColor(0x1D, 0x1F, 0x22)
MUTED = RGBColor(0x55, 0x58, 0x5E)
FAINT = RGBColor(0x8A, 0x8D, 0x93)


# ---------- image helpers ----------

def crop_image(path: str, crop, boxes=(), size: int = 640) -> io.BytesIO:
    """Crop (normalised x,y,w,h) and draw frame-coordinate boxes [(bbox, colour, width)]."""
    im = load_resized(path, tuple(crop) if crop else None).convert("RGB")
    im.thumbnail((size, size))
    d = ImageDraw.Draw(im)
    cx, cy, cw, ch = crop or (0, 0, 1, 1)
    W, H = im.size
    for bbox, colour, width in boxes:
        bx, by, bw, bh = bbox
        x0, y0 = (bx - cx) / cw * W, (by - cy) / ch * H
        x1, y1 = x0 + bw / cw * W, y0 + bh / ch * H
        d.rectangle([x0, y0, x1, y1], outline=colour, width=width)
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=82, optimize=True)
    buf.seek(0)
    return buf


# ---------- docx helpers ----------

def shade(cell, hex_fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def set_cell_borders(table, colour="D6D3CA") -> None:
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), colour)
        borders.append(el)
    tblPr.append(borders)


def para(doc, text="", size=10.5, bold=False, colour=INK, space_after=6, italic=False, align=None):
    p = doc.add_paragraph()
    if text:
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold = bold
        r.italic = italic
        r.font.color.rgb = colour
    p.paragraph_format.space_after = Pt(space_after)
    if align:
        p.alignment = align
    return p


def eyebrow(doc, text):
    p = para(doc, text.upper(), size=8, colour=FAINT, space_after=2)
    for r in p.runs:
        r.font.name = "Consolas"
    return p


def heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.color.rgb = INK
        r.font.name = "Arial"
    return h


def add_table(doc, rows, col_widths_cm, header=True, font_size=9.5, zebra=False):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    set_cell_borders(t)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.width = Cm(col_widths_cm[j])
            cell.text = ""
            p = cell.paragraphs[0]
            r = p.add_run(str(val))
            r.font.size = Pt(font_size)
            if header and i == 0:
                r.bold = True
                r.font.color.rgb = MUTED
                shade(cell, "E9E7E0")
            elif j == 0:
                r.font.color.rgb = MUTED
            p.paragraph_format.space_after = Pt(2)
    return t


def image_row(doc, images, widths_cm, labels):
    """One table row of images with small labels beneath."""
    t = doc.add_table(rows=2, cols=len(images))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for j, (buf, w, lab) in enumerate(zip(images, widths_cm, labels)):
        c = t.cell(0, j)
        c.width = Cm(w)
        c.paragraphs[0].add_run().add_picture(buf, width=Cm(w - 0.2))
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        lc = t.cell(1, j)
        lc.width = Cm(w)
        r = lc.paragraphs[0].add_run(lab.upper())
        r.font.size = Pt(7)
        r.font.color.rgb = FAINT
        r.font.name = "Consolas"
        lc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    return t


def pct(x):
    return f"{x * 100:.0f}%"


# ---------- data (mirrors build_report.main) ----------

def main() -> None:
    opus_sum, opus = br.load_run(br.OPUS_DIR)
    haiku_sum, haiku = br.load_run(br.HAIKU_DIR) if br.HAIKU_DIR.exists() else (None, {})
    fo = br.floor_only(opus_sum, opus)
    fh = br.floor_only(haiku_sum, haiku) if haiku_sum else None
    cs, crows = br.load_cctv(br.CCTV_CROPS)
    fs, frows = br.load_cctv(br.CCTV_FULL)
    if not cs:
        sys.exit("no fixed-camera results; run scripts/run_cctv_eval.py --mode crops --three-frames first")

    n_pairs = cs["pairs"]
    cpos = [r for r in crows if r["has_spill"]]
    cneg = [r for r in crows if not r["has_spill"]]
    seqs = sorted({r["seq"] for r in crows})
    ccf = cs["confusion"]
    fcf = fs["confusion"] if fs else {"tp": 0, "fn": 0, "fp": 0, "tn": 0}
    HANDHELD = {"walkway_vert"}
    mounted = [r for r in crows if r["seq"] not in HANDHELD]
    m_tp = sum(1 for r in mounted if r["has_spill"] and r["pred"])
    m_fn = sum(1 for r in mounted if r["has_spill"] and not r["pred"])
    m_fp = sum(1 for r in mounted if not r["has_spill"] and r["pred"])
    m_tn = sum(1 for r in mounted if not r["has_spill"] and not r["pred"])
    m_acc = (m_tp + m_tn) / len(mounted)
    calls_per = cs["model_calls"] / n_pairs
    c_cost = cs["this_run"]["est_cost_usd"] / cs["this_run"]["pairs"]
    f_cost = fs["this_run"]["est_cost_usd"] / fs["this_run"]["pairs"] if fs else 0
    c_lat = sum(r["latency_s"] for r in crows) / len(crows)
    f_lat = sum(r["latency_s"] for r in frows) / len(frows) if frows else 0

    # showcase selection, same rules as the HTML report
    tps = [r for r in cpos if r["pred"]]
    misses = [r for r in cpos if not r["pred"]]
    show = []
    for s in seqs:
        seen = set()
        for r in sorted((t for t in tps if t["seq"] == s), key=lambda t: -t["confidence"]):
            if r["kind"] in seen:
                continue
            seen.add(r["kind"])
            show.append((r, f"{s} camera"))
            if len(seen) == 4:
                break
    show += [(r, f"{r['seq']} camera, a miss") for r in misses[:2]]
    show += [(r, f"{r['seq']} camera, people walking through") for r in [n for n in cneg if n["candidates"] and n["calls"]][:2]]

    water = [r for r in cpos if r["kind"] == "water"]
    water_found = [r for r in water if r["pred"]]
    wshow = []
    for s in seqs:
        best = sorted((r for r in water_found if r["seq"] == s), key=lambda r: -r["confidence"])[:1]
        wshow += [(r, f"{s} camera, water") for r in best]
    smallest = sorted(water_found, key=lambda r: r["spill_width_frac"])[:1]
    wshow += [(r, f"{r['seq']} camera, the smallest water spill found") for r in smallest if r["id"] not in {x["id"] for x, _ in wshow}]
    wshow += [(r, f"{r['seq']} camera, a water spill it missed") for r in [r for r in water if not r["pred"]][:1]]

    # ---------- document ----------
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    content_w = 17.0

    eyebrow(doc, "Proof of concept · Mall security cameras · Opus 5.5 vision · 21 October 2026")
    heading(doc, "Spill detection on the mall's own cameras", 0)
    para(doc, "The mall already has fixed cameras over every corridor and atrium. This proof of concept uses them to spot a liquid "
              "spill within seconds of it happening and to send the cleaning robot only to confirmed spills, so the robot does not "
              "patrol. No training data was collected and no model was trained: a general vision model judges what changed on the "
              "floor against a written description of a spill, and every answer is kept as a label for a smaller on-site model later.",
         size=11.5, colour=MUTED, space_after=10)

    tiles = [
        (pct(m_acc), "accurate on mounted cameras", f"{len(mounted)} before-and-after checks, 4 cameras"),
        (pct(m_tp / (m_tp + m_fn)), "of spills found", f"{m_tp} of {m_tp + m_fn} rendered spills on real floors"),
        (str(m_fp), "false alarms", f"in {m_fp + m_tn} real frames of people and trolleys"),
        (f"${c_cost:.2f}", "per checked frame", "frames with nothing persistent cost nothing"),
    ]
    t = doc.add_table(rows=1, cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell_borders(t)
    for j, (big, label, src) in enumerate(tiles):
        c = t.cell(0, j)
        c.width = Cm(content_w / 4)
        shade(c, "FBFAF7")
        p = c.paragraphs[0]
        r = p.add_run(big)
        r.bold = True
        r.font.size = Pt(22)
        p2 = c.add_paragraph()
        r2 = p2.add_run(label)
        r2.bold = True
        r2.font.size = Pt(9)
        p3 = c.add_paragraph()
        r3 = p3.add_run(src)
        r3.font.size = Pt(8)
        r3.font.color.rgb = FAINT
    para(doc, "", space_after=4)

    eyebrow(doc, "How it works")
    heading(doc, "Reference, persistent change, verdict", 1)
    para(doc, "Because the cameras never move, the server keeps a reference picture of each empty floor and only asks the model about "
              "patches that changed and stayed changed: a spill persists, a person walks on. Each such patch is cut out of the frames "
              "at full size and sent with the reference and a frame two seconds later. A frame with no persistent change costs nothing. "
              "The robot is dispatched only to a confirmed spill.", colour=MUTED)
    steps = [
        ("1 · Reference. Remember the empty floor.", "A median over a few seconds of footage removes passers-by and leaves the clean floor, per camera. Refreshed slowly so lighting changes do not pile up."),
        ("2 · Persistent change. Find what appeared and stayed.", "Difference against the reference, kept only where the change is still there two seconds later. People, trolleys and their shadows drop out; spills, dropped bags and signs remain."),
        ("3 · Ask about the patch. Three frames, one crop.", "Claude sees the same crop from the reference, the current frame and a frame two seconds later, and answers \"has a liquid spill appeared on the floor?\" with a fixed JSON shape. The third frame is how it tells a spill, which stays, from a shadow, which walks away."),
    ]
    for title, body in steps:
        p = para(doc, "", space_after=4)
        r = p.add_run(title + " ")
        r.bold = True
        r.font.size = Pt(10.5)
        r2 = p.add_run(body)
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = MUTED

    eyebrow(doc, "Results")
    para(doc, f"Across all {n_pairs} pairs from {len(seqs)} cameras: {pct(cs['accuracy'])} accuracy, {ccf['tp']} of {ccf['tp'] + ccf['fn']} spills found, "
              f"{ccf['fp']} false alarms in {ccf['tn'] + ccf['fp']} real frames, and {cs['spill_located_correctly']} found spills placed at the right spot. "
              f"On the four mounted cameras alone ({len(mounted)} pairs): {pct(m_acc)} accuracy, {m_tp} of {m_tp + m_fn} spills found, "
              f"{m_fp} false alarms in {m_fp + m_tn} real frames. The handheld walkway clip accounts for every false alarm.", colour=MUTED)

    def pair_block(r, caption):
        gt_c = None
        if r["gt_bbox"]:
            gt_c = (r["gt_bbox"][0] + r["gt_bbox"][2] / 2, r["gt_bbox"][1] + r["gt_bbox"][3] / 2)
        calls = [c for c in r["calls"] if "spill" in c]
        call = next((c for c in calls if c["spill"]), None)
        if call is None and gt_c:
            call = next((c for c in calls if c.get("crop") and br._contains(c["crop"], gt_c)), None)
        if call is None and calls:
            call = calls[0]
        crop = call["crop"] if call and call.get("crop") else [0, 0, 1, 1]
        boxes = []
        if r["gt_bbox"]:
            boxes.append((r["gt_bbox"], GREEN, 2))
        if call and call.get("spill") and call.get("bbox_frame"):
            boxes.append((call["bbox_frame"], AMBER, 3))
        three = bool(r.get("later")) and r["later"] != r["cur"]
        imgs = [crop_image(r["ref"], crop), crop_image(r["cur"], crop, boxes)]
        labels = ["reference", "current"]
        if three:
            imgs.append(crop_image(r["later"], crop))
            labels.append("2 s later")
        w = min((content_w - 0.4) / len(imgs), 5.6)  # two-image rows stay the same height as three-image rows
        image_row(doc, imgs, [w] * len(imgs), labels)
        pred = r["pred"]
        conf = call["confidence"] if call else 0.0
        reason = call["reasoning"] if call else r["reasoning"]
        outcome = ("Found" if pred else "Missed") if r["has_spill"] else ("False alarm" if pred else "Correctly clear")
        truth = f"rendered {r['kind']} spill, {r['spill_width_frac']*100:.0f}% of frame width" if r["has_spill"] else "no spill, real footage"
        p = para(doc, "", space_after=2)
        rr = p.add_run(("SPILL" if pred else "CLEAR") + f"   confidence {conf:.2f}   {outcome}")
        rr.bold = True
        rr.font.size = Pt(9)
        rr.font.color.rgb = RGBColor(0x5A, 0x40, 0x00) if pred else RGBColor(0x0F, 0x4D, 0x31)
        para(doc, reason, size=9.5, space_after=1)
        para(doc, f"{caption} · truth: {truth}", size=8, colour=FAINT, space_after=10)

    eyebrow(doc, "What the model saw")
    para(doc, "Each example is exactly what the model saw: the same crop from the reference, the current frame and two seconds later. "
              "The amber box is the model's answer, the thin green box is the true spill.", colour=MUTED)
    for r, cap in show:
        pair_block(r, cap)

    eyebrow(doc, "Clear liquids")
    heading(doc, "Water is the hard case, and it is found", 2)
    para(doc, f"Water has no colour of its own. On a floor it shows only as a darkening of the surface, a soft outline and a slight sheen, "
              f"which is also how a shadow or a polished patch looks. The model found {len(water_found)} of {len(water)} water spills across the "
              f"mall cameras, down to one {min(r['spill_width_frac'] for r in water_found)*100:.0f}% of the frame wide. "
              + "Coloured liquids for comparison: " + ", ".join(
                  f"{k} {sum(1 for r in cpos if r['kind']==k and r['pred'])} of {sum(1 for r in cpos if r['kind']==k)}"
                  for k in ("coffee", "cola", "juice", "oil")) + ".", colour=MUTED)
    for r, cap in wshow:
        pair_block(r, cap)

    heading(doc, "The pipeline against a whole-frame check", 2)
    add_table(doc, [
        ["", f"Pipeline: persistent crops, three frames ({n_pairs} pairs)", f"Whole frame, two frames ({fs['pairs'] if fs else 0} pairs, every third)"],
        ["Spills found", f"{ccf['tp']} of {ccf['tp'] + ccf['fn']}", f"{fcf['tp']} of {fcf['tp'] + fcf['fn']}"],
        ["False alarms", f"{ccf['fp']} of {ccf['tn'] + ccf['fp']}", f"{fcf['fp']} of {fcf['tn'] + fcf['fp']}"],
        ["Spills smaller than 7% of frame width found",
         f"{sum(1 for r in cpos if r['spill_width_frac'] < 0.07 and r['pred'])} of {sum(1 for r in cpos if r['spill_width_frac'] < 0.07)}",
         f"{sum(1 for r in frows if r['has_spill'] and r['spill_width_frac'] < 0.07 and r['pred'])} of {sum(1 for r in frows if r['has_spill'] and r['spill_width_frac'] < 0.07)}"],
        ["Model calls per frame", f"{calls_per:.1f}", "1"],
        ["Cost per checked frame", f"${c_cost:.3f}", f"${f_cost:.3f}"],
        ["Latency per frame", f"{c_lat:.1f} s", f"{f_lat:.1f} s"],
    ], [6.0, 5.5, 5.5])
    para(doc, "", space_after=4)

    eyebrow(doc, "Cost at scale")
    heading(doc, "What this costs across a mall", 1)
    para(doc, f"The camera server does the cheap work for free: it keeps the reference, finds changes and drops the ones that walk away. "
              f"Only persistent patches reach the model, and each check costs about ${c_cost:.3f} at the {calls_per:.1f} crops per check seen in "
              f"the busy test clips. One camera, a 12-hour day:", colour=MUTED)
    add_table(doc, [
        ["One camera", "Checks per day", "Per day", "Per month"],
        ["Quiet floor: one check every 4 minutes", "180", f"${180 * c_cost:,.2f}", f"${180 * 30 * c_cost:,.0f}"],
        ["Busy corridor: one check every minute", "720", f"${720 * c_cost:,.2f}", f"${720 * 30 * c_cost:,.0f}"],
    ], [7.0, 3.3, 3.3, 3.4])
    para(doc, "The per-check cost was measured on crowded test clips where people standing still add crops; an empty floor costs less per check. "
              "Once enough labelled patches have accumulated, a small on-site model takes the easy decisions and the cloud model becomes a second opinion.",
         size=9.5, colour=FAINT, space_after=8)

    eyebrow(doc, "Next")
    heading(doc, "Roadmap", 1)
    add_table(doc, [
        ["When", "What"],
        ["Done", "Fixed-camera pipeline. Reference, persistent change, before-and-after crops. Tested on five mall cameras with rendered spills: no false alarms on the mounted cameras."],
        ["Next", "One real mall camera, real spills. A few hours of footage from a single camera in any mall, with water, coffee and oil spilled on the floor in view. This replaces the rendered positives and gives the true recall number and the real check rate."],
        ["Then", "Hard negatives at scale. Polished stone, reflections, cleaning carts, wet-floor signs, dropped bags. Measure the false-alarm rate per camera per day, which is what decides whether staff trust the alerts."],
        ["Later", "Camera-server deployment. Reference and persistence run on the camera server at no cost; only persistent patches go to the model. Robot dispatch on confirmed spills. An on-site model trained from the accumulated labels where the cloud call is not wanted."],
    ], [2.5, 14.5])

    # ---------- close camera ----------
    doc.add_page_break()
    eyebrow(doc, "Also tested")
    heading(doc, "A close camera on the robot itself", 1)
    para(doc, "Before the camera question was settled, the same model was tested on single close-range frames, the view a camera on the "
              "robot would have. It is kept here because it is the fallback if a mall's cameras cannot be used, and because it shows how "
              "the model behaves on a public spill dataset with real liquids.", colour=MUTED)
    para(doc, "One or two frames per second are sampled from the camera, near-identical frames are dropped, and each frame is sent to "
              "Claude Opus 5.5 with a written brief of what counts as a spill and what does not. The answer is a yes/no, a confidence, "
              "a liquid type, a box and a one-line reason, in a fixed JSON shape, and every answer is kept with its frame as a training label.",
         colour=MUTED)

    heading(doc, "What the model saw, and what it said", 2)
    para(doc, "Real close-range frames from the public dataset. The amber box is where the model says the spill is. The text under each "
              "frame is the model's own one-line reason, unedited.", colour=MUTED)
    cards = []
    for prefix, group, caption in br.SHOWCASE:
        r = br.find(opus, prefix)
        if r and not r.get("error"):
            cards.append((r, caption))
    for i in range(0, len(cards), 2):
        pair = cards[i:i + 2]
        t = doc.add_table(rows=2, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, (r, caption) in enumerate(pair):
            x = r["result"]
            boxes = [(x["bbox"], AMBER, 3)] if x["bbox"] else []
            c = t.cell(0, j)
            c.width = Cm(8.3)
            c.paragraphs[0].add_run().add_picture(crop_image(r["path"], None, boxes, 560), width=Cm(8.0))
            lc = t.cell(1, j)
            lc.width = Cm(8.3)
            p = lc.paragraphs[0]
            rr = p.add_run(("SPILL" if x["spill_detected"] else "CLEAR") + f"   confidence {x['confidence']:.2f}")
            rr.bold = True
            rr.font.size = Pt(9)
            rr.font.color.rgb = RGBColor(0x5A, 0x40, 0x00) if x["spill_detected"] else RGBColor(0x0F, 0x4D, 0x31)
            p2 = lc.add_paragraph()
            r2 = p2.add_run(x["reasoning"])
            r2.font.size = Pt(9)
            p3 = lc.add_paragraph()
            r3 = p3.add_run(caption)
            r3.font.size = Pt(8)
            r3.font.color.rgb = FAINT
        para(doc, "", space_after=4)

    if haiku_sum and fh:
        heading(doc, "Why Opus 5.5 and not the cheaper model", 2)
        para(doc, f"Claude Haiku 5.5 costs about a thirtieth as much. On the same {len(opus)} frames its headline score is close, but it is less "
                  f"consistent: it ignored the floor-only rule on {fh['non_floor'] - fh['rule_ok']} of the {fh['non_floor']} mislabelled frames, "
                  f"gave up on several abstract splash images, and missed the one genuinely hard floor spill in the set.", colour=MUTED)
        hr = haiku_sum["this_run"]
        cost_per = opus_sum["this_run"]["est_cost_usd"] / opus_sum["this_run"]["images"]
        hlat = sum(r["latency_s"] for r in haiku.values()) / len(haiku)
        lat = sum(r["latency_s"] for r in opus.values()) / len(opus)
        add_table(doc, [
            ["", "Claude Opus 5.5", "Claude Haiku 5.5"],
            ["Floor-frame accuracy", pct(fo["accuracy"]), pct(fh["accuracy"])],
            ["Floor spills missed", f"{fo['fn']} of {fo['tp'] + fo['fn']}", f"{fh['fn']} of {fh['tp'] + fh['fn']}"],
            ["False alarms", f"{fo['fp']} of {fo['tn'] + fo['fp']}", f"{fh['fp']} of {fh['tn'] + fh['fp']}"],
            ["Followed the floor-only rule", f"{fo['rule_ok']} of {fo['non_floor']}", f"{fh['rule_ok']} of {fh['non_floor']}"],
            [f"Cost, {len(opus)} frames", f"${cost_per * len(opus):.2f}", f"${hr['est_cost_usd'] / hr['images'] * len(haiku):.3f}"],
            ["Mean latency", f"{lat:.1f} s", f"{hlat:.1f} s"],
        ], [6.0, 5.5, 5.5])
        para(doc, "Recommendation: Opus 5.5 for the proof of concept and for labelling anything subtle. Haiku may later take the easy frames "
                  "in bulk, with Opus re-checking anything it marks below 0.8 confidence.", size=9.5, colour=FAINT)

    para(doc, "", space_after=6)
    for line in (
        "Model: claude-opus-5-5, effort low, structured JSON output.",
        "Fixed-camera footage: five mall clips from Pixabay and Pexels (Pixabay Content License, Pexels License), real floors with rendered spills; no AI image generation was used anywhere in this work.",
        f"Close-camera data: Spill_Detection v4 by slipp on Roboflow Universe, CC BY 4.0. Validation split, {len(opus)} frames, random seed 0.",
        "Generated from data/results by scripts/build_docx.py.",
    ):
        para(doc, line, size=8, colour=FAINT, space_after=1)

    doc.save(OUT)
    print(f"wrote {OUT} ({OUT.stat().st_size / 1024 / 1024:.1f} MB)")


if __name__ == "__main__":
    main()
